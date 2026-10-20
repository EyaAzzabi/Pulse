"""Transaction import from Excel, CSV or a Google Sheets link, with a
user-chosen column mapping. Re-importing the same rows is a no-op: each row's
import_reference is a hash of its mapped values."""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from urllib.parse import parse_qs, urlparse

import httpx
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.financial import Transaction

MAX_ROWS = 5000
FIELDS = ("date", "amount", "direction", "category", "description")
REQUIRED = ("date", "amount")
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d")

SYNONYMS = {
    "date": ["date", "transaction date", "booking date", "value date", "posted", "posting date", "operation date"],
    "amount": ["amount", "montant", "value", "total", "sum", "amount (eur)", "amount eur"],
    "direction": ["direction", "type", "debit/credit", "dc", "sens", "in/out", "flow"],
    "category": ["category", "catégorie", "categorie", "account", "class"],
    "description": ["description", "label", "libellé", "libelle", "memo", "details", "narration", "reference", "counterparty"],
}

IN_WORDS = {"in", "credit", "cr", "income", "deposit", "receipt", "inflow", "+", "crédit", "entrée", "entree"}
OUT_WORDS = {"out", "debit", "dr", "expense", "withdrawal", "payment", "outflow", "-", "débit", "sortie"}


class ImportSourceError(Exception):
    pass


@dataclass
class Table:
    headers: list[str]
    rows: list[list]


@dataclass
class ImportOutcome:
    imported: int = 0
    skipped_duplicates: int = 0
    errors: list[dict] = field(default_factory=list)


def read_table(filename: str, content: bytes) -> Table:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return _read_xlsx(content)
    if name.endswith(".csv") or not name:
        return _read_csv(content)
    raise ImportSourceError("Only .xlsx and .csv files are supported.")


def _read_xlsx(content: bytes) -> Table:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportSourceError("This file couldn't be opened as an Excel workbook.")
    ws = wb.worksheets[0]
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return _to_table(all_rows)


def _read_csv(content: bytes) -> Table:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return _to_table([row for row in csv.reader(io.StringIO(text), dialect)])


def _to_table(all_rows: list[list]) -> Table:
    rows = [r for r in all_rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise ImportSourceError("The file is empty.")
    headers = [str(h).strip() if h is not None else f"Column {i + 1}" for i, h in enumerate(rows[0])]
    body = rows[1 : MAX_ROWS + 1]
    width = len(headers)
    return Table(headers=headers, rows=[(list(r) + [None] * width)[:width] for r in body])


def google_sheet_export_url(url: str) -> str:
    """Builds the CSV export URL ourselves rather than fetching whatever URL
    the user pasted, so this endpoint can't be used to reach other hosts."""
    parsed = urlparse(url.strip())
    if parsed.scheme not in ("http", "https") or parsed.hostname != "docs.google.com":
        raise ImportSourceError("Paste a link that starts with https://docs.google.com/spreadsheets/.")
    match = re.match(r"^/spreadsheets/d/([a-zA-Z0-9_-]+)", parsed.path)
    if not match:
        raise ImportSourceError("This doesn't look like a Google Sheets link.")
    gid = parse_qs(parsed.query).get("gid", [""])[0]
    if not gid:
        frag = re.search(r"gid=(\d+)", parsed.fragment)
        gid = frag.group(1) if frag else "0"
    if not re.fullmatch(r"\d+", gid):
        gid = "0"
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=csv&gid={gid}"


def fetch_google_sheet(url: str) -> Table:
    export_url = google_sheet_export_url(url)
    try:
        response = httpx.get(export_url, follow_redirects=True, timeout=15.0)
    except httpx.HTTPError:
        raise ImportSourceError("Couldn't reach Google Sheets. Try again in a moment.")
    content_type = response.headers.get("content-type", "")
    if response.status_code != 200 or "text/csv" not in content_type:
        raise ImportSourceError(
            "PULSE can't read this sheet. In Google Sheets, click Share and set access to "
            "'Anyone with the link can view'."
        )
    return _read_csv(response.content)


def suggest_mapping(headers: list[str]) -> dict[str, str | None]:
    lowered = {h: h.lower().strip() for h in headers}
    mapping: dict[str, str | None] = {}
    used: set[str] = set()
    for fld in FIELDS:
        choice = None
        for syn in SYNONYMS[fld]:
            choice = next((h for h, l in lowered.items() if l == syn and h not in used), None)
            if choice:
                break
        if not choice:
            choice = next(
                (h for h, l in lowered.items() if h not in used and any(syn in l for syn in SYNONYMS[fld])), None
            )
        if choice:
            used.add(choice)
        mapping[fld] = choice
    return mapping


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    raw = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def _parse_amount(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    raw = str(value).strip()
    negative = raw.startswith("(") and raw.endswith(")")
    raw = re.sub(r"[^\d,.\-]", "", raw)
    if "," in raw and "." in raw:
        raw = raw.replace(",", "")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        amount = Decimal(raw)
    except InvalidOperation:
        return None
    return -abs(amount) if negative else amount


def _parse_direction(value) -> str | None:
    if value is None:
        return None
    return "IN" if str(value).strip().lower() in IN_WORDS else "OUT" if str(value).strip().lower() in OUT_WORDS else None


def import_table(db: Session, organization_id: str, table: Table, mapping: dict[str, str | None]) -> ImportOutcome:
    missing = [f for f in REQUIRED if not mapping.get(f)]
    if missing:
        raise ImportSourceError(f"Choose a column for: {', '.join(missing)}.")
    index = {}
    for fld, header in mapping.items():
        if header:
            if header not in table.headers:
                raise ImportSourceError(f"Column '{header}' isn't in this file.")
            index[fld] = table.headers.index(header)

    existing = {
        r[0]
        for r in db.query(Transaction.import_reference).filter(Transaction.organization_id == organization_id).all()
        if r[0]
    }
    outcome = ImportOutcome()
    for n, row in enumerate(table.rows, start=2):
        cell = lambda f: row[index[f]] if f in index else None  # noqa: E731
        tx_date = _parse_date(cell("date"))
        amount = _parse_amount(cell("amount"))
        if tx_date is None:
            outcome.errors.append({"row": n, "reason": f"Unreadable date: '{cell('date')}'"})
            continue
        if amount is None or amount == 0:
            outcome.errors.append({"row": n, "reason": f"Unreadable amount: '{cell('amount')}'"})
            continue

        if "direction" in index:
            direction = _parse_direction(cell("direction"))
            if direction is None:
                outcome.errors.append({"row": n, "reason": f"Unknown direction: '{cell('direction')}'"})
                continue
        else:
            direction = "IN" if amount > 0 else "OUT"

        category = str(cell("category") or "").strip() or "uncategorized"
        description = str(cell("description") or "").strip()[:255]
        abs_amount = abs(amount).quantize(Decimal("0.01"))
        ref = hashlib.sha256(
            f"{tx_date.isoformat()}|{abs_amount}|{direction}|{category}|{description}".encode()
        ).hexdigest()[:32]
        if ref in existing:
            outcome.skipped_duplicates += 1
            continue
        db.add(
            Transaction(
                organization_id=organization_id,
                date=tx_date,
                amount=abs_amount,
                currency="EUR",
                category=category[:64],
                direction=direction,
                source="file_import",
                status="posted",
                description=description,
                import_reference=ref,
            )
        )
        existing.add(ref)
        outcome.imported += 1
    db.commit()
    return outcome


def preview(table: Table) -> dict:
    return {
        "headers": table.headers,
        "sample_rows": [[_display(c) for c in r] for r in table.rows[:5]],
        "total_rows": len(table.rows),
        "suggested_mapping": suggest_mapping(table.headers),
    }


def _display(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else str(value)
