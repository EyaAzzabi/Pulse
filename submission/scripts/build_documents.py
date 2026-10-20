"""Generate evidence-based reports and a coordinated editable pitch deck."""
import json
import math
from pathlib import Path
import textwrap
from xml.sax.saxutils import escape

from PIL import Image as PILImage
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image, KeepTogether
from reportlab.graphics.shapes import Drawing, Rect, String, Line, Polygon
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "submission"
E = OUT / "execution_evidence"
S = OUT / "screenshots"
M = json.loads((E / "demo_metrics.json").read_text())
NAVY = "10233E"; DARK = "0B172B"; CYAN = "5DDBEF"; TEAL = "52C7AA"; MUTED = "9BAFC8"; INK = "183353"
FONT = "SegoeUI"; BOLD = "SegoeUIBold"
pdfmetrics.registerFont(TTFont(FONT, "C:/Windows/Fonts/segoeui.ttf"))
pdfmetrics.registerFont(TTFont(BOLD, "C:/Windows/Fonts/segoeuib.ttf"))
money = lambda v: f"EUR {float(v):,.2f}"
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="PulseBody", fontName=FONT, fontSize=9.5, leading=15, spaceAfter=9, textColor=colors.HexColor("#334D69")))
styles.add(ParagraphStyle(name="PulseH1", fontName=BOLD, fontSize=24, leading=29, spaceAfter=18, textColor=colors.HexColor("#" + NAVY)))
styles.add(ParagraphStyle(name="PulseH2", fontName=BOLD, fontSize=14, leading=20, spaceAfter=10, spaceBefore=10, textColor=colors.HexColor("#" + INK)))
styles.add(ParagraphStyle(name="PulseSmall", fontName=FONT, fontSize=8, leading=12, textColor=colors.HexColor("#607B99")))


def para(text, style="PulseBody"):
    return Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])


def table(headers, rows, widths=None):
    data = [[para(x, "PulseSmall") for x in headers]] + [[para(x, "PulseSmall") for x in row] for row in rows]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F0FA")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8), ("LINEBELOW", (0, 0), (-1, -1), .4, colors.HexColor("#DDE6F1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFD")])]))
    return t


def figure(name, caption, width=475):
    path = S / name
    if not path.exists(): return [para("Screenshot unavailable: " + name)]
    with PILImage.open(path) as im: ratio = im.height / im.width
    return [Image(str(path), width=width, height=width * ratio), Spacer(1, 6), para(caption, "PulseSmall"), Spacer(1, 15)]


def diagram(labels, footer):
    d = Drawing(485, 125)
    w = 108; gap = 14
    for i, label in enumerate(labels):
        x = i * (w + gap)
        d.add(Rect(x, 42, w, 57, fillColor=colors.HexColor("#" + NAVY), strokeColor=None, rx=7, ry=7))
        lines = textwrap.wrap(label, 18)
        for j, line in enumerate(lines): d.add(String(x + 10, 78 - j * 14, line, fontName=FONT, fontSize=9, fillColor=colors.white))
        if i < len(labels)-1:
            d.add(Line(x+w+2, 70, x+w+gap-3, 70, strokeColor=colors.HexColor("#208B9F"), strokeWidth=2))
    d.add(String(0, 15, footer, fontName=FONT, fontSize=8, fillColor=colors.HexColor("#637D99")))
    return d


def page_frame(c, doc):
    w, h = doc.pagesize
    c.setFillColor(colors.HexColor("#" + NAVY)); c.rect(0, h-15, w, 15, fill=1, stroke=0)
    c.setFont(FONT, 8); c.setFillColor(colors.HexColor("#7890AA"))
    c.drawString(48, 26, "PULSE AI  /  FinTechathon 2026  /  Topic E")
    c.drawRightString(w-48, 26, str(doc.page))


def report(path, title, subtitle, sections):
    doc = SimpleDocTemplate(str(path), pagesize=(595.28, 841.89), rightMargin=48, leftMargin=48, topMargin=45, bottomMargin=48,
        title=title, author="PULSE AI team")
    story = [Spacer(1, 35), para("PULSE AI", "PulseH1"), para(title, "PulseH1"), para(subtitle), Spacer(1, 20),
        para("Shenzhen International FinTech Competition 2026\nInternational Track — Topic E: SME Finance Copilot"),
        Spacer(1, 25), para("Evidence date: 20 October 2026. All example business records are synthetic. No performance or adoption metrics are assumed. Supabase/PostgreSQL is the team's shared deployment database; the evidence run uses isolated SQLite.", "PulseSmall"), PageBreak()]
    for idx, (heading, content) in enumerate(sections):
        story.append(para(heading, "PulseH1"))
        for item in content:
            if isinstance(item, str): story.append(para(item))
            elif isinstance(item, list): story.extend(item)
            else: story.append(item)
        if idx != len(sections)-1: story.append(PageBreak())
    doc.build(story, onFirstPage=page_frame, onLaterPages=page_frame)


technical = [
("01 / Executive overview", [
    "PULSE AI is a decision-support financial copilot for SMEs. A deterministic financial engine transforms transactions, invoices and recurring obligations into liquidity analysis and a baseline forecast. A Financial Digital Twin recomputes that same forecast under explicit alternative assumptions. The copilot adds explanations and tool orchestration; human approval controls consequential actions.",
    "The central design principle is separation: financial arithmetic belongs to deterministic code; language generation is optional; simulations do not overwrite financial records; action proposals do not execute themselves.",
    table(["Verified synthetic example", "Actual result"], [["Current cash", money(M['summary']['current_cash'])], ["First forecast shortfall", M['shortfall_date']], ["Baseline lowest cash / days negative", money(M['baseline_lowest']) + ' / ' + str(M['baseline_days_below_zero'])], ["Supplier terms +15 days: lowest / days negative", money(M['rescue_lowest']) + ' / ' + str(M['rescue_days_below_zero'])]], [280, 195]),
    "These are demonstration outputs, not a forecast for a real SME and not evidence of predictive accuracy."]),
("02 / The SME problem and solution", [
    "A business can show a positive historical cash-flow average while still facing a near-term timing gap. Owners often lack a dedicated CFO, and spreadsheets do not automatically connect overdue collections, scheduled payroll, supplier bills and alternative interventions.",
    "PULSE helps the owner ask four practical questions: What is my current position? When could cash fall below zero? Which assumption or obligation causes the dip? What changes if I act differently?",
    "In the synthetic case, trailing 30-day net flow is " + money(M['summary']['net_cash_flow_30d']) + ". Historical burn is zero, so the burn-based runway is undefined rather than infinite. The scheduled-flow forecast nevertheless identifies a shortfall in " + str(M['baseline_runway']) + " days. This distinction is intentional.",
    figure("04_financial_dashboard.png", "Figure 1. Real dashboard connected to the isolated synthetic demo backend.")]),
("03 / System architecture", [
    diagram(["React / TypeScript browser", "FastAPI authorization + APIs", "Financial engines + tools", "Supabase PostgreSQL"], "Shared deployment architecture. Evidence backend: isolated SQLite with the same ORM and engines."),
    "Browser requests carry a JWT bearer token. FastAPI resolves the authenticated user, requires account approval for business APIs, then checks organization membership and minimum role. The scoped financial input loader supplies records to the engines. Results are persisted as forecast runs, scenarios, proposals, conversations, notifications and audit records.",
    "Supabase is used as the team's shared PostgreSQL deployment database, through SQLAlchemy's PostgreSQL connection. This is not a claim that Supabase Auth or PostgreSQL row-level-security policies are implemented. PULSE authentication and tenant checks remain application-level controls. The current local .env and submission fixtures use SQLite; a Supabase connectivity check is not evidenced here.",
    "Email uses authenticated n8n webhooks and SMTP, with direct SMTP fallback for OTP. No bank or payment API is present. The submission action demonstration disables both outbound channels and honestly records QUEUED."]),
("04 / Technology and data model", [
    table(["Layer", "Actual implementation"], [["Frontend", "React 19, TypeScript, Vite, Tailwind CSS, React Query, Recharts, Framer Motion, Axios, React Router"], ["Backend", "FastAPI, Pydantic, SQLAlchemy, Alembic, Python Decimal"], ["Database", "Supabase/PostgreSQL shared deployment; SQLite isolated evidence fixture"], ["Identity", "bcrypt password hashing; JWT bearer tokens; HMAC-protected email OTP; platform account review"], ["AI", "OpenAI-compatible tool calling, Groq/OpenAI configuration, deterministic fallback"], ["Automation", "n8n 2.44.4 local harness, SMTP STARTTLS/implicit TLS, authenticated completion receipts"], ["Import", "CSV, Excel via openpyxl, supported Google Sheets URLs and column mapping"]], [100,375]),
    "The ORM models users and memberships; organizations, business profiles and onboarding progress; transactions, invoices and recurring obligations; forecast runs/points; simulation scenarios/results; action proposals and recommendations; conversations/messages; notifications, audit logs, OTPs and account reviews. The complete column-level schema is included in execution_evidence/database_schema.json.",
    "The requirements file is the reproducible dependency input. The captured runtime version manifest is separate and does not assume that installed packages exactly match every pinned requirement."]),
("05 / Financial algorithms", [
    "All currency arithmetic uses Decimal and quantization to two decimal places. Current liquidity is opening cash plus recorded inflows minus recorded outflows through the as-of date. The opening balance is defined before the first recorded transaction. Onboarding adjusts the stored opening balance when the user supplies a current balance after importing history.",
    "Trailing net flow = sum(inflows) − sum(outflows) over the requested inclusive dates. Monthly burn = max(outflows − inflows, 0) / (lookback_days / 30). Burn-based runway = floor(current cash / (monthly burn / 30)); it is null when burn is not positive.",
    "Outstanding and overdue balances use invoice amount minus amount_paid, excluding paid invoices. Concentration is the largest customer's outstanding receivables divided by total outstanding receivables. Expense trend compares recent and preceding outflow windows.",
    "Risk rules flag negative cash, burn-based runway below 30 or 60 days, overdue receivables, and customer concentration of at least 50%. These are transparent rules, not a trained credit-risk score.",
    "Independent evidence replays raw demo transactions to obtain " + money(M['summary']['current_cash']) + ", validates every one of the 90 daily forecast recurrences, checks the no-change simulation against baseline, and verifies the exact EUR 8,000 day-one inventory delta. Invoice snapshots remain unchanged after simulations and queued execution.",
    "Limits: calculations do not perform foreign-exchange conversion or accounting reconciliation. The current liquidity loop does not explicitly filter transaction status. Inclusive comparison windows can share a boundary date; these definitions need review before accounting-grade use."]),
("06 / Cash-flow forecast methodology", [
    "For each projected day: cash[d] = cash[d−1] + known[d] + uncertain[d]. Known flow includes remaining invoice balances scheduled on their due date and modeled recurring obligations. Monthly obligations match day-of-month; weekly obligations repeat every seven days.",
    "Uncertain flow is a flat trailing 90-day daily net average after excluding categories represented by recurring obligations. It is included only when there are at least 30 historical transactions. Confidence labels are high through day 14, medium through day 45, then low; insufficient history yields low confidence throughout.",
    "These confidence labels are heuristic categories, not calibrated probabilities or prediction intervals. There is no seasonality model, macroeconomic forecast or trained time-series model. Unpaid overdue invoices are not automatically assumed collectible in the baseline. Invoice cash receipts and historical sales averages can overlap, so double-counting beyond recurring-category exclusion remains a modeling concern.",
    figure("05_cashflow_forecast.png", "Figure 2. Actual 90-day forecast for the synthetic demo. First negative date: " + str(M['shortfall_date']))]),
("07 / Financial Digital Twin", [
    "Scenario assumptions operate on copied inputs. Supported changes include customer/supplier delays, revenue/expense shocks, one-time inventory purchases, financing proceeds/repayments and one-time cash injections. The internal tool layer additionally supports customer-specific delay and a stated overdue-collection assumption.",
    "Both baseline and simulated paths use run_forecast. Comparison includes lowest cash point/date, days below zero and days until first shortfall. API impact_vs_baseline measures ending cash difference, not improvement of the minimum. A timing-only intervention may improve the cash minimum while leaving ending cash unchanged.",
    "Supplier terms +15 days raises the actual simulated minimum by " + money(float(M['rescue_lowest'])-float(M['baseline_lowest'])) + ", from " + money(M['baseline_lowest']) + " to " + money(M['rescue_lowest']) + ". This removes negative-cash days in this particular synthetic scenario. It is conditional on the assumed extension and is not a guarantee of supplier consent.",
    figure("09_scenario_comparison.png", "Figure 3. Real supplier-extension simulation and baseline comparison.")]),
("08 / Specialist agents and tools", [
    diagram(["Question + scoped context", "Orchestrator: intent / tool loop", "Financial / forecast / strategy tools", "Explanation + proposal"], "Risk Guardian: role/scope checks and prompt constraints. Workflow execution remains outside the LLM loop."),
    "The specialists are logical roles in one application, not separately deployed autonomous services. The tool registry includes get_financial_summary, forecast_cash, simulate_scenario, get_rescue_plans, list_invoices and propose_invoice_reminder. Tool results carry computed values and evidence.",
    "With a provider key, the orchestrator permits up to six model/tool rounds and dispatches only registered tools. Unknown tools are rejected. Reminder drafting requires at least ACCOUNTANT. The model does not receive a direct execution tool. With no provider, keyword-based deterministic responses and traces remain available.",
    "The demo and screenshots use deterministic mode. LLM tool-loop behavior is covered by repository tests with mocked provider responses, not by an actual external model inference in this evidence run. Generated language is not completely mechanically verified against tool outputs; avoid claiming that every possible LLM statement is guaranteed correct.",
    figure("07_ai_copilot.png", "Figure 4. Genuine deterministic copilot response and trace.")]),
("09 / Retrieval and data grounding", [
    "RAG is not implemented. There is no embedding pipeline, vector database or semantic document retrieval index. Grounding is through structured financial input loading and registered deterministic tools.",
    "The Excel/CSV/Google Sheets importer is a tabular ingestion feature, not retrieval-augmented generation. It offers column mapping, preview and duplicate detection. File size and row limits exist, but streaming upload limits, compressed workbook resource limits and adversarial input handling require hardening.",
    "For a future RAG extension, keep document retrieval read-only, enforce tenant scope before retrieval, attach provenance and distinguish retrieved claims from engine-computed figures. This is a roadmap item, not submission functionality."]),
("10 / Automation and human approval", [
    "The five n8n exports cover signup verification, cash-risk email, a financial digest, approved reminder dispatch and a recovery callback. Header-authenticated webhooks consume the request body and return explicit completion receipts. The local digest copy is on demand, while the source export also contains a schedule. Enabling recurring delivery requires explicit operational configuration.",
    "PULSE creates action proposals, requires an authorized human approval, then checks approval and organization scope again at execution. The recovery workflow calls this guarded endpoint instead of bypassing it. SENT requires a confirmed submission; absence of a configured dispatcher yields QUEUED.",
    "In this package's demo, no customer email is sent. The recorded sequence proves HTTP 409 before approval, UI approval, QUEUED after execution, and HTTP 409 on sequential replay. An earlier local n8n integration run confirmed all five workflows with real SMTP; its sanitized receipt summary is included separately with its provenance.",
    figure("10_human_approval.png", "Figure 5. Actual pending approval screen in the synthetic workspace."),
    figure("12_n8n_workflow.png", "Figure 6. Real authenticated n8n editor, if captured; no credential fields are exposed.")]),
("11 / Authentication, tenancy and audit", [
    "New signup verifies an email OTP, stores a bcrypt password hash and creates a PENDING account. Platform approval is a separate privilege from organization ownership. Pending/rejected/suspended accounts can read their status but cannot use business APIs. Account-review updates compare the expected current status and record a review event atomically with the change.",
    "Organization endpoints check membership. Protected operations require minimum roles. Conversation lookup additionally binds the conversation to both the organization and requesting user. Administrative access does not automatically grant access to tenant finances.",
    "OTP codes are randomly generated, HMAC-hashed, short-lived, single-use, and have attempt/resend limits. Failed delivery removes an unusable code. SMTP supports STARTTLS and implicit TLS. These protections do not replace global rate limiting, deployment TLS or operational monitoring.",
    "The browser stores bearer tokens in localStorage, exposing them to any successful XSS. Audit tables are application history, not cryptographically immutable evidence. For full threat modeling, priority fixes and control status, see the separate security assessment."]),
("12 / Verification and reproduction", [
    "Fresh backend verification: 86 tests passed; the raw terminal output is provided. The frontend production build exited zero; it reports a large-chunk warning, not a functional failure. The automated evidence browser recorded no JavaScript page errors in the exercised routes.",
    "Financial validation assertions and raw API request/response pairs are included. Twelve genuine screenshots and a real Chromium recording document the available pages. No network responses were mocked for these submission screenshots. All records shown belong to the isolated synthetic demo.",
    "Reproduce: install backend requirements and frontend lockfile dependencies; run submission/scripts/demo_environment.py; start its serve mode on port 8002; run Vite on port 5175 with VITE_API_BASE_URL=http://127.0.0.1:8002; run capture_evidence.py. Its private state is generated under ignored .runtime and is excluded from the ZIP.",
    "Shared deployment: every user's frontend must target the same backend, and that backend must use the same Supabase DATABASE_URL. Merely copying source code to two laptops does not share account state. Apply migrations to the intended shared database and configure SMTP/n8n on that backend.",
    "The source commit and archive hash are documented separately. Uploading or generating this package does not constitute official competition submission."]),
]

controls = [
["JWT + bcrypt", "IMPLEMENTED", "Signed expiring bearer tokens; bcrypt hashes; no automatic refresh/revocation scheme proven."],
["Email OTP", "IMPLEMENTED", "HMAC hashes, expiry, attempts and resend cooldown; successful mailbox delivery separately confirmed in local work."],
["Account approval", "IMPLEMENTED", "Platform admin review; non-approved business requests blocked."],
["Tenant checks", "IMPLEMENTED", "Membership and role checks at application level; not Supabase RLS."],
["Tool allowlist", "IMPLEMENTED", "Registered tools and role-checked reminder proposals; no direct execution tool."],
["Output correctness", "PARTIAL", "Deterministic calculations and prompt grounding; LLM narrative not exhaustively verified."],
["Replay protection", "PARTIAL", "Status rejects sequential replay. Concurrent races and provider-side idempotency need work."],
["Auditability", "PARTIAL", "Database action/account logs exist. No append-only external storage or tamper-evidence."],
["Input limits", "PARTIAL", "Pydantic validation and import limits; full request/resource caps are incomplete."],
["Secret handling", "PARTIAL", "Local .env/runtime ignored; deployment secret manager/rotation not evidenced."],
["Global throttling / MFA", "NOT IMPLEMENTED", "No general login/IP rate limiter or MFA in inspected code."],
["Deployment assurance", "NOT VERIFIED", "No penetration test, Supabase RLS validation, backup restoration drill or dependency CVE audit performed."],
]
security = [
("01 / Scope and threat model", ["This is an engineering self-assessment, not a security certification or penetration-test report. Scope: the inspected PULSE source, the synthetic backend/browser run, and observed local n8n configuration. Supabase is the selected shared PostgreSQL deployment, but its live connectivity and database policies were not verified in this local session.",
    table(["Asset", "Threat / trust boundary"], [["Accounts and bearer tokens", "Credential guessing, token theft through XSS, account takeover; browser-to-API boundary."], ["Tenant financial records", "Cross-organization reads/writes, excessive role permissions; API-to-database boundary."], ["Financial decisions", "Misleading forecasts, LLM unsupported claims, malicious imported data or prompt content."], ["Outbound workflows", "Unauthenticated webhooks, duplicate dispatch, wrong recipients, forged success receipts."], ["Operational secrets", "Leaked .env, signing keys, SMTP credentials, database URLs or n8n encryption keys."], ["Availability", "Large uploads, unbounded forecasts, expensive parsing or provider delays."]], [145,330]),
    "Threat actors include unauthenticated callers, malicious tenant users, compromised browsers, adversarial content supplied to the copilot/importer, and operators with excessive database access. No claim is made that all attacks have been tested."]),
("02 / Control status matrix", [table(["Control", "Status", "Evidence and boundary"], controls, [105,100,270])]),
("03 / Permission model", [table(["Tier", "Observed authority"], [["Unauthenticated", "Signup OTP/register/login endpoints; no business APIs."], ["Pending / rejected / suspended", "Authenticated status visibility; business access denied."], ["VIEWER", "Read tenant financial data and permitted analytical simulations."], ["ACCOUNTANT", "Draft supported action proposals and import financial data, subject to route role checks."], ["FINANCE_MANAGER", "Approve and execute supported actions; lower-role capabilities."], ["OWNER", "Highest organization role; organization profile/management operations."], ["Platform administrator", "Account approvals, organization metadata, system configuration summaries, account-review log. No automatic tenant finance bypass."]], [140,335]),
    "Role membership is checked server-side. Hiding a UI route is not the security boundary. The action and account-review tests exercise unauthorized access and state transitions."]),
("04 / Authentication and secrets", ["Passwords are hashed with bcrypt. JWTs contain a subject and expiry and are signed using configured settings. Account active/approval state is re-read from the database, so a status suspension blocks existing tokens on subsequent API calls.",
    "Email OTP is an email-possession check, not identity verification or MFA. Per-code attempt limits and per-address resend cooldown exist; IP-wide anti-abuse, distributed request throttling and enumeration resistance are incomplete. Email-existence responses can reveal whether an account is registered.",
    "JWTs in localStorage are exposed if an attacker achieves XSS. There is no proven refresh-token family, central session revocation or stolen-token detection. Change signing keys and review session invalidation procedures before production.",
    "Actual .env, credential files, runtime databases and n8n state are excluded from this package. A scanner compares publishable files against local SMTP/JWT/webhook secrets without printing them. .env.example is a placeholder template. Defaults and example secrets must never become production signing keys."]),
("05 / Tenant isolation and Supabase", ["The primary isolation mechanism is require_membership plus require_role and organization-filtered queries. User-specific conversation access is also checked. Tests cover cross-organization access attempts. These are application controls and depend on every new endpoint applying them correctly.",
    "Supabase/PostgreSQL is the team-selected shared database. The code does not establish Supabase Auth integration or a validated RLS policy set. A privileged database connection could bypass database-level isolation unless operationally restricted. Use a dedicated least-privilege application role, encrypted connections, network controls and reviewed policies.",
    "Two laptops each using sqlite:///./pulse.db do not share registrations. The admin and every signup frontend must point to one shared backend/database deployment. The current local .env remained SQLite during capture; the evidence fixture was deliberately isolated. No claim of a passed Supabase connection or migration test is made."]),
("06 / AI, prompt injection and financial risk", ["Unknown tools are denied. Tool inputs are tied to the authenticated organization, and reminder drafting enforces a minimum role. The orchestration loop is bounded in rounds, but not every model output is mechanically verified. A trace saying Risk Guardian review is not proof of a full semantic audit.",
    "Treat user prompts, customer names, invoice descriptions, imported cells and retrieved provider messages as untrusted. Prompt instructions request grounded numbers, but an LLM can still produce unsupported language, select a poor scenario or misinterpret ambiguity. Do not describe prompt rules as an absolute hallucination blocker.",
    "No RAG pipeline is implemented. Proposed future retrieval must enforce tenant scope before indexing/retrieval and label sources. Risks from cross-tenant retrieval are therefore roadmap concerns, not currently tested RAG controls.",
    "Forecasts use assumptions and historical averages, not guarantees. Confidence labels are heuristic. Missing/duplicated data, invoice-versus-history overlap, month-end recurrences, inconsistent statuses and currencies can materially distort outputs. Add reconciliation, model validation and user-visible assumptions before production reliance."]),
("07 / Actions, webhooks and audit", ["Approval and execution are separate endpoints. Execution requires FINANCE_MANAGER or OWNER, organization membership and APPROVED state. Sequential repeated execution is rejected. The synthetic demonstration proves the before-approval and replay denials without contacting a customer.",
    "Known weakness: the current execution route reads status and dispatches before an atomic committed claim, so concurrent calls could race. A provider timeout may mean email was accepted but acknowledgment was lost; a fallback could duplicate delivery. Introduce atomic compare-and-set claims, durable outbox entries, provider idempotency keys and receipt reconciliation.",
    "n8n webhooks use Header Auth and explicit receipts; production deployment needs HTTPS. Local HTTP is used only on loopback. Execution payload storage is disabled in the supplied local exports to avoid retaining OTP content. Any real mail recipient must be validated; a customer name is not an address.",
    "Audit logs capture approval/execution/account-review events. Database operators can still modify tables. Stronger evidence requires restricted write privileges, append-only storage or signed external logs, retention policies and monitored access."]),
("08 / Input validation and availability", ["Pydantic constrains registration fields and several scenario/import inputs. File import enforces a 5 MB check after reading the upload and a 5,000-row processing limit. This is not a streaming memory guard; compressed spreadsheets can consume more resources when expanded.",
    "Forecast and some scenario horizon entry points need stricter bounds. Analytical simulation may be available to low-privilege members and still create persisted scenario records. Define quotas, request timeouts, bounded payloads and lifecycle cleanup.",
    "Google Sheets fetching uses supported source parsing; redirects and allowlists should be reviewed for SSRF exposure. No live SSRF exploit or adversarial workbook test was performed. Do not claim these risks have been eliminated.",
    "No comprehensive dependency CVE scan, external penetration test, load test, backup restore drill or incident-response exercise was performed. The test count demonstrates regression coverage, not proof of total security."]),
("09 / Mitigation priorities", [table(["Priority", "Action", "Acceptance evidence"], [["P0", "Centralize backend on the shared Supabase deployment; use unique production secrets and TLS.", "Verified shared signup/admin flow; migration log; secret/config review."], ["P0", "Atomically claim actions and add durable idempotency/outbox processing.", "Concurrent execution tests and timeout reconciliation without duplicate sends."], ["P0", "Add IP/account rate limiting and strict compute/upload bounds.", "Abuse tests and resource-budget checks."], ["P1", "Harden browser sessions, CSP and XSS defenses; add MFA for admins.", "Session threat review and browser security tests."], ["P1", "Strengthen tenant database privileges/RLS where appropriate and audit new routes.", "Database policy tests and endpoint isolation suite."], ["P1", "Validate LLM claims and monetary/date outputs against tool evidence.", "Adversarial prompt suite and structured response checks."], ["P2", "Calibrate forecast uncertainty, improve reconciliation and operational monitoring.", "Documented validation dataset, restore drills and alert exercises."]], [45,230,200]),
    "These are recommendations and planned work, not completed controls."]),
("10 / Evidence and conclusion", ["The fresh regression run reports 86 tests passing. Actual financial checks, API responses, screenshot manifest and browser error results are included. The production frontend build succeeds with a large JavaScript chunk warning. The submission browser flow reports zero page errors on the exercised routes.",
    "The no-key copilot path is verified here. Live LLM inference, Supabase policy behavior, security penetration testing and statistical forecast accuracy are not verified. Existing mocked provider tests must not be presented as an external model run.",
    "PULSE is a functional prototype with meaningful permission, approval and audit controls. It should be presented as financial decision support with explicit assumptions, and hardened before handling production financial decisions at scale."])
]


class Slide:
    def __init__(self, prs, pdf, index, label):
        self.slide = prs.slides.add_slide(prs.slide_layouts[6]); self.pdf=pdf
        self.rect(0,0,13.333,7.5,DARK)
        self.text(.55,.32,"PULSE AI",14,CYAN,bold=True)
        self.text(9,.35,"FINTECHATHON 2026  /  TOPIC E",8,MUTED,width=3.8)
        self.line(.55,7.08,12.78,7.08,"243750",.8)
        self.text(.55,7.18,label.upper(),8,MUTED,width=10)
        self.text(12.4,7.15,f"{index:02d}",11,CYAN,width=.4)

    def rect(self,x,y,w,h,color,radius=False):
        sh=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
        self.pdf.setFillColor(colors.HexColor('#'+color));self.pdf.rect(x*72,(7.5-y-h)*72,w*72,h*72,fill=1,stroke=0)

    def line(self,x1,y1,x2,y2,color,width=1):
        from pptx.enum.shapes import MSO_CONNECTOR
        sh=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
        self.pdf.setStrokeColor(colors.HexColor('#'+color));self.pdf.setLineWidth(width);self.pdf.line(x1*72,(7.5-y1)*72,x2*72,(7.5-y2)*72)

    def text(self,x,y,text,size=20,color="FFFFFF",width=11.8,bold=False):
        font=BOLD if bold else FONT
        lines=[]
        for raw in str(text).split('\n'):
            line=''
            for word in raw.split():
                trial=(line+' '+word).strip()
                if pdfmetrics.stringWidth(trial,font,size)>width*72 and line:lines.append(line);line=word
                else:line=trial
            lines.append(line)
        height=max(.25,len(lines)*size*1.25/72+.12)
        sh=self.slide.shapes.add_textbox(Inches(x),Inches(y),Inches(width),Inches(height));tf=sh.text_frame;tf.clear();tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
        for i,ln in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=ln;p.font.name="Segoe UI";p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0);p.line_spacing=1.15
            self.pdf.setFont(font,size);self.pdf.setFillColor(colors.HexColor('#'+color));self.pdf.drawString(x*72,(7.5-y)*72-size*.85-i*size*1.15,ln)
        return height

    def image(self,name,x,y,w,h):
        path=S/name
        with PILImage.open(path) as im:ratio=im.width/im.height
        width=min(w,h*ratio);height=width/ratio
        self.slide.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(width),height=Inches(height))
        self.pdf.drawImage(ImageReader(str(path)),x*72,(7.5-y-height)*72,width*72,height*72)

    def note(self,text):self.slide.notes_slide.notes_text_frame.text=text
    def end(self):self.pdf.showPage()


def deck():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    pdf=canvas.Canvas(str(OUT/'PULSE_Presentation.pdf'),pagesize=(960,540));pdf.setTitle('PULSE AI — Financial foresight for SMEs')
    titles=['The Invisible Financial Crisis','Why SMEs Struggle','Introducing PULSE AI','The Financial Digital Twin','Predicting Financial Risk','From Risk to Rescue','Agentic AI Architecture','Live Demo and Execution Evidence','Security and Human Control','Impact and Future Vision']
    def new(i):return Slide(prs,pdf,i,titles[i-1])
    s=new(1);s.text(.65,1.1,'Your business has a heartbeat.',38,width=9.9,bold=True);s.text(.65,2.15,'See the cash crisis\nbefore it arrives.',45,CYAN,width=10.5,bold=True)
    s.text(.7,4.3,'PULSE AI',27,bold=True);s.text(.7,4.9,'The Financial Autopilot for SMEs\nDecision support. Human control.',19,MUTED,width=7)
    s.rect(9,4.45,3.45,1.75,NAVY);s.text(9.25,4.68,'18 DAYS',33,TEAL,width=3,bold=True);s.text(9.25,5.38,'to the demo forecast shortfall',13,MUTED,width=2.9)
    s.note('Timing: 45 seconds. Open with the gap between financial history and the next obligation. In our synthetic example, cash is positive today but the engine forecasts a negative position in eighteen days. These are real outputs from our demonstration data, not a claim about a real customer or prediction accuracy. PULSE is a financial copilot, not a service that autonomously moves money. We use the term autopilot to describe continuous analytical support while the owner keeps control.');s.end()
    s=new(2);s.text(.65,1.05,titles[1],35,bold=True);s.text(.7,2,'Positive history.\nA negative cash date.',29,CYAN,width=5.2,bold=True)
    for y,a,b in [(3.6,'REACTIVE','Spreadsheets explain yesterday.'),(4.5,'FRAGMENTED','One late invoice changes the picture.'),(5.4,'UNSUPPORTED','Owners need evidence, not guesses.')]:s.text(.75,y,a,11,TEAL,bold=True);s.text(.75,y+.26,b,16,MUTED,width=5.1)
    s.image('06_liquidity_risk_alert.png',6.2,2,6.5,4.6);s.note('Timing: 50 seconds. SMEs may have accounting data without an integrated view of cash timing. Our demonstration has positive trailing net flow and no positive historical burn rate, yet the forecast sees a timing shortfall. The product connects invoices, transactions and recurring obligations so an owner can inspect the underlying reason. We do not claim market adoption, customer savings or forecast accuracy measurements. The problem and design are illustrated by our synthetic, reproducible business case.');s.end()
    s=new(3);s.text(.65,1.05,titles[2],35,bold=True);s.text(.7,2,'Understand.\nAnticipate.\nSimulate.\nDecide.',31,CYAN,width=4.3,bold=True);s.text(.7,5.3,'Real calculations first.\nLanguage generation is optional.',16,MUTED,width=4.5);s.image('04_financial_dashboard.png',5.1,2,7.65,4.85);s.note('Timing: 50 seconds. PULSE is a working React and FastAPI application backed by deterministic Python engines. Supabase PostgreSQL is our shared deployment database; this captured demonstration deliberately uses a separate SQLite fixture, so no private business data appears. The dashboard combines current liquidity, upcoming events, risks and candidate interventions. The no-key copilot remains functional because the numbers do not depend on an LLM. This separation is the core product architecture, not a claim that generated language can never make mistakes.');s.end()
    s=new(4);s.text(.65,1.05,titles[3],35,bold=True);s.text(.7,2,'One business.\nMore possible futures.',29,CYAN,width=4.1,bold=True);s.text(.7,3.65,'Payment timing\nInventory purchases\nRevenue and expense shocks\nFinancing assumptions',17,MUTED,width=4.1);s.text(.7,5.75,'Original financial records stay intact.',13,TEAL,width=4.3);s.image('08_digital_twin_simulator.png',5.1,2,7.65,4.8);s.note('Timing: 50 seconds. The Digital Twin copies inputs and runs the same forecast engine for both baseline and alternative. We can test payment delays, an inventory purchase, revenue or expense changes and financing assumptions. The EUR 8,000 purchase in our evidence shifts day-one cash by exactly EUR 8,000. We independently validate that result and check that the original invoice records remain unchanged. Scenarios are stored for comparison, but the financial records they model are not rewritten. Financing examples are illustrative terms, not real lender offers.');s.end()
    s=new(5);s.text(.65,1.05,titles[4],35,bold=True);s.text(.7,1.9,'Known commitments + uncertain daily flow',20,MUTED)
    x0,y0,w,h=.8,3.2,8.8,2.7;vals=[float(p['projected_cash']) for p in M['forecast_points']];lo=min(vals)-1500;hi=max(vals)+1500
    for i in range(1,len(vals)):
        s.line(x0+(i-1)/(len(vals)-1)*w,y0+h-(vals[i-1]-lo)/(hi-lo)*h,x0+i/(len(vals)-1)*w,y0+h-(vals[i]-lo)/(hi-lo)*h,CYAN,2)
    zero=y0+h-(0-lo)/(hi-lo)*h;s.line(x0,zero,x0+w,zero,'52627A',1);s.text(.8,6.1,'Actual 90-day engine output · synthetic business · estimates, not guarantees',11,MUTED)
    s.text(10,2.8,money(M['baseline_lowest']),25,TEAL,width=2.7,bold=True);s.text(10,3.55,'lowest projected cash',14,MUTED,width=2.6);s.text(10,4.65,str(M['baseline_days_below_zero'])+' DAYS',29,CYAN,width=2.6,bold=True);s.text(10,5.35,'below zero in baseline',14,MUTED,width=2.6);s.note('Timing: 50 seconds. This chart is drawn from the actual forecast JSON, not a fabricated interface. Each day equals the previous balance plus known scheduled flow and an explicitly uncertain historical average. Recurring categories are excluded from that average to avoid one category-level double count. Confidence labels are heuristic and not statistically calibrated probabilities. The cash minimum is negative EUR 2,308, with seven negative-cash days. We also disclose model limitations such as month-end recurrence and possible overlap between invoice collections and average historical revenue.');s.end()
    s=new(6);s.text(.65,1.05,titles[5],35,bold=True);s.text(.7,1.95,'What if suppliers grant 15 more days?',24,CYAN,width=11.5)
    for x,label,value,color in [(.75,'BASELINE',money(M['baseline_lowest']),'F3A9A0'),(5,'SIMULATED',money(M['rescue_lowest']),TEAL),(9.2,'NEGATIVE DAYS','7 → 0',CYAN)]:s.rect(x,2.8,3.4,1.65,NAVY);s.text(x+.2,3.05,label,10,MUTED);s.text(x+.2,3.5,value,27,color,width=3.05,bold=True)
    s.text(.8,5,'EUR 4,372 improvement in the lowest cash point.',24,TEAL,width=11.5,bold=True);s.text(.8,5.7,'A conditional scenario, not a guarantee. Collection and credit alternatives are also ranked by engine results.',17,MUTED,width=11.5);s.note('Timing: 50 seconds. We choose a specific intervention and state its assumption: suppliers agree to a fifteen-day extension. The simulated cash minimum rises from negative EUR 2,308 to positive EUR 2,064, an improvement of EUR 4,372. Negative-cash days fall from seven to zero in this case. The engine also ranks collection and illustrative credit strategies. These are conditional outcomes, not guaranteed rescue promises. Importantly, ending-cash impact and minimum-cash improvement are separate metrics: shifting timing can improve the minimum without changing the ending balance.');s.end()
    s=new(7);s.text(.65,1.05,titles[6],35,bold=True);s.text(.7,1.9,'Specialist roles. Registered tools. Scoped data.',21,MUTED)
    for i,(a,b) in enumerate([('Financial Analyst','Liquidity and invoices'),('Forecast Analyst','Dates and risk drivers'),('Strategic Planner','Simulations and rescues'),('Risk Guardian','Tool scope and permissions')]):
        x=.7+i*3.15;s.rect(x,3,2.95,1.65,NAVY);s.text(x+.18,3.3,a,18,CYAN,width=2.6,bold=True);s.text(x+.18,4,b,12,MUTED,width=2.55)
    s.text(.8,5.15,'Orchestrator → deterministic tools → explanation / proposal',22,TEAL,width=11.8);s.text(.8,6,'No RAG claim. No direct payment tool. LLM tool calling is optional.',14,MUTED);s.note('Timing: 55 seconds. These are logical specialist roles in one service, not four independently deployed AI services. The registry exposes financial summary, forecasting, simulation, rescue-plan, invoice and reminder-proposal tools. With an LLM key, a bounded tool loop chooses registered calls. Without a key, the deterministic path still works. The Risk Guardian enforces scope and permissions, but we do not claim a complete semantic verifier of all generated text. There is no RAG implementation or vector database. No tool directly sends money, and proposed reminders still need approval.');s.end()
    s=new(8);s.text(.65,1.05,titles[7],35,bold=True);s.image('07_ai_copilot.png',.65,2,6.15,3.85);s.image('11_execution_history.png',6.95,2,5.8,3.8);s.text(.8,6,'86 tests passed  ·  Real browser capture  ·  No customer dispatch in this demo',17,TEAL,width=11.9);s.note('Timing: 120 seconds for the included genuine browser video, or about 60 seconds for the screenshot walkthrough. Show the dashboard, risk alert, copilot, purchase scenario, supplier extension, approval and queued execution. This prototype run uses synthetic data and deterministic copilot mode. It demonstrates an actual blocked execution before approval, human approval through the UI, queued action when no external channel is configured, and a blocked sequential replay. QUEUED must never be described as SENT. The separate n8n integration evidence confirms real SMTP workflows from the local integration session.');s.end()
    s=new(9);s.text(.65,1.05,titles[8],35,bold=True);s.text(.7,2,'Propose → Review → Approve → Execute → Audit',24,CYAN,width=12)
    for i,(title,body) in enumerate([('Identity','JWT, bcrypt, email OTP'),('Scope','Organization and role checks'),('Control','Approval before execution')]):
        x=.75+i*4.15;s.rect(x,3,3.9,1.7,NAVY);s.text(x+.2,3.28,title,22,TEAL,bold=True,width=3.5);s.text(x+.2,3.95,body,16,MUTED,width=3.4)
    s.text(.8,5.3,'Honest boundaries',22,bold=True);s.text(.8,5.85,'Next: stronger concurrent idempotency, rate limits, session hardening and validated deployment policies.',17,MUTED,width=11.8);s.note('Timing: 50 seconds. Identity and tenant controls are implemented server-side. A platform administrator reviews account access but does not automatically gain tenant financial access. Action approval and audit history are real. We are also explicit about gaps: browser tokens are in localStorage, audit tables are not tamper-proof, and concurrency-safe execution and external provider idempotency need strengthening. Supabase is the selected shared database, but database-level RLS is not established by these tests. This is a self-assessment, not a certification or penetration-test result.');s.end()
    s=new(10);s.text(.65,1.1,'Better foresight.\nMore deliberate decisions.',41,CYAN,width=11.5,bold=True);s.text(.75,3.35,'IMPLEMENTED NOW',11,TEAL,bold=True);s.text(.75,3.8,'Deterministic engine + Digital Twin + optional copilot + human-approved workflows',24,width=11.7)
    s.text(.75,5.15,'NEXT CHAPTER',11,TEAL,bold=True);s.text(.75,5.55,'Shared Supabase deployment · Bank feeds · Calibrated uncertainty · Production hardening',18,MUTED,width=11.6);s.note('Timing: 50 seconds. Close on the concrete differentiation: the same engine powers insight, forecast, what-if comparison and evidence-based interventions. This reduces the gap between a business question and an inspectable decision. We do not claim a measured reduction in business failures or customer savings. Supabase PostgreSQL is our shared deployment direction; banking integrations, calibrated uncertainty and deeper operational hardening are future work. The competition package contains actual tests, API outputs, source instructions, screenshots and a genuine demonstration video so judges can reproduce the implementation.');s.end()
    prs.save(OUT/'PULSE_Presentation.pptx');pdf.save()


if __name__ == '__main__':
    report(OUT/'PULSE_Technical_Documentation.pdf','Technical Documentation','Verified architecture, financial methods, implementation evidence and limitations.',technical)
    report(OUT/'PULSE_Security_Self_Assessment.pdf','Security Self-Assessment','Implemented controls, partial safeguards, known design weaknesses and mitigation priorities.',security)
    deck()
    print('Generated technical PDF, security PDF, 10-slide PPTX and coordinated presentation PDF.')
