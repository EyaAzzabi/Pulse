"""Builds the PULSE AI submission PDFs and slide deck from real project facts and screenshots.

Run from the repository root: python submission/build_docs.py
Requires: playwright (chromium) and python-pptx.
"""
import base64
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SUB = Path(__file__).resolve().parent
SHOTS = SUB / "screenshots"
NAVY, CYAN, TEAL = "#0b1f3a", "#22d3ee", "#14b8a6"


def img(name, caption=""):
    data = base64.b64encode((SHOTS / f"{name}.png").read_bytes()).decode()
    return f'<figure><img src="data:image/png;base64,{data}"/><figcaption>{caption}</figcaption></figure>'


CSS = f"""
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600&family=Poppins:wght@600;700&display=swap');
body{{font-family:Inter,Arial,sans-serif;color:#1e293b;font-size:10.5pt;line-height:1.5;margin:0}}
h1,h2,h3{{font-family:Poppins,Arial,sans-serif;color:{NAVY}}}
h1{{font-size:26pt;margin:0 0 4px}} h2{{font-size:15pt;border-bottom:2px solid {TEAL};padding-bottom:4px;margin-top:26px;page-break-after:avoid}}
h3{{font-size:11.5pt;margin-bottom:4px}}
.cover{{background:{NAVY};color:white;padding:48px 40px;border-radius:14px;margin-bottom:20px}}
.cover h1{{color:white}} .cover p{{color:{CYAN};margin:4px 0}}
table{{border-collapse:collapse;width:100%;margin:8px 0;font-size:9.5pt;page-break-inside:avoid}}
th{{background:{NAVY};color:white;text-align:left;padding:6px}} td{{border-bottom:1px solid #e2e8f0;padding:6px;vertical-align:top}}
code,pre{{font-family:Consolas,monospace;font-size:9pt;background:#f1f5f9;border-radius:4px}} pre{{padding:10px;white-space:pre-wrap}}
figure{{margin:10px 0;page-break-inside:avoid}} img{{width:100%;border:1px solid #cbd5e1;border-radius:8px}}
figcaption{{font-size:8.5pt;color:#64748b;text-align:center}}
.ok{{color:#047857;font-weight:600}} .part{{color:#b45309;font-weight:600}} .plan{{color:#64748b;font-weight:600}}
.note{{background:#ecfeff;border-left:4px solid {CYAN};padding:8px 12px;margin:10px 0}}
"""

ARCH = """<pre>
 Browser (React 19 + Vite + Tailwind)            n8n (optional, signed webhooks)
   Landing / Onboarding / Dashboard / Copilot       verification code, invoice reminder,
   Forecast / Simulator / Decisions / Admin         cash risk alert, digest, recovery plan
              | HTTPS + JWT bearer                          ^  HMAC/header secret
              v                                             |
 FastAPI backend  ------------------------------------------+
   api/        auth, onboarding, imports, financial, forecast, simulations,
               recommendations, actions, copilot, notifications, admin
   agents/     orchestrator (OpenAI tool calling, deterministic fallback) + tools.py
   forecasting/ deterministic cash forecast engine (Decimal)
   simulations/ Financial Digital Twin (scenario recalculation)
   services/   rescue plans, alerts, notifications, OTP, email, tabular import
              | SQLAlchemy 2 + Alembic
              v
 Supabase Postgres (session pooler)   |  SQLite for the automated test suite
</pre>"""

FLOW = """<pre>
 Financial data -> Summary + risk -> 90-day forecast -> shortfall detected
   -> Digital Twin scenarios -> ranked rescue plans -> agent proposes action
   -> PENDING_APPROVAL -> human approves (role checked) -> execute
   -> QUEUED (no webhook) or dispatched via n8n -> audit log + notification
</pre>"""


def tech_doc():
    return f"""<html><head><meta charset="utf-8"><style>{CSS}</style></head><body>
<div class="cover"><h1>PULSE AI</h1><p>Technical Documentation</p>
<p>Shenzhen International FinTech Competition 2026 (FinTechathon), International Track, Topic E</p>
<p>Repository: github.com/EyaAzzabi/Pulse (branch main)</p></div>

<h2>1. Overview</h2>
<p>PULSE AI is a cash-flow copilot for small and medium enterprises. It turns transactions, invoices and recurring
obligations into a deterministic forecast, detects upcoming shortfalls, simulates rescue strategies in a Financial
Digital Twin and lets AI agents propose actions that only run after a human with the right role approves them.</p>
<p>Every figure shown to the user comes from the deterministic engines. The language model chooses tools and phrases
answers; it never invents numbers.</p>

<h2>2. Problem</h2>
<p>SMEs often see a cash crisis only when it arrives: late customer payments, clustered supplier bills and tax dates
combine into a negative balance weeks later. Spreadsheets do not show the date of the shortfall or which action helps most.</p>

<h2>3. Solution</h2>
<ul><li><b>Onboarding</b> with email verification code, company profile, role, and data connection (Excel/CSV upload with column mapping, Google Sheets link, Northstar Studio sample company, or quick estimates).</li>
<li><b>Dashboard</b>: current cash, 30-day net flow, lowest point in 60 days, overdue receivables, insight card, upcoming 30 days, overdue aging, what changed.</li>
<li><b>Forecast</b>: day-by-day cash projection with known and uncertain components.</li>
<li><b>Simulator</b>: what-if scenarios (customer delay, supplier delay, revenue/expense change, financing, injections).</li>
<li><b>Copilot</b>: multi-agent chat with conversation history, voice input and read-aloud, and a visible agent trace.</li>
<li><b>Decision Center</b>: propose, approve, execute, with audit log.</li>
<li><b>Notifications</b> and <b>Admin console</b> (users, organizations, account review, service status, audit).</li></ul>

<h2>4. System architecture</h2>{ARCH}
{img("04_dashboard","Dashboard, Northstar Studio, captured from the running app")}

<h2>5. Core scenario flow</h2>{FLOW}

<h2>6. Data model</h2>
<table><tr><th>Entity</th><th>Purpose</th></tr>
<tr><td>User, Organization, Membership</td><td>Tenancy and role per organization (OWNER, FINANCE_MANAGER, ACCOUNTANT, VIEWER); is_admin, account_status, email_verified</td></tr>
<tr><td>Transaction, Invoice, RecurringObligation</td><td>Financial inputs; money stored as Decimal</td></tr>
<tr><td>ForecastRun / points</td><td>Stored forecasts with starting cash and risk</td></tr>
<tr><td>Scenario / SimulationResult</td><td>Digital Twin runs and assumptions</td></tr>
<tr><td>Action, Recommendation</td><td>Proposed actions with status PENDING_APPROVAL, APPROVED, QUEUED, SENT, FAILED</td></tr>
<tr><td>Conversation, Message</td><td>Copilot memory with stored agent trace</td></tr>
<tr><td>Notification, AuditLog, AccountReview, EmailOtp</td><td>Alerts, traceability, admin review, hashed verification codes</td></tr></table>

<h2>7. Core algorithms</h2>
<h3>7.1 Current liquidity and summary</h3>
<p>Current cash = opening balance + sum of all recorded transactions up to today. Runway, 30-day net flow and overdue
receivables are computed from the same inputs.</p>
<h3>7.2 Deterministic forecast</h3>
<p>Starting from current liquidity, each future day adds <i>known</i> flows (open invoices on their due date, recurring
obligations on their schedule) and an <i>uncertain</i> baseline derived from the historical daily average of categories
not already modeled as recurring (to avoid double counting). Output: daily balance, lowest point, first shortfall date,
risk level and reasons.</p>
{img("06_forecast","Forecast page")}
<h3>7.3 Financial Digital Twin</h3>
<p>The simulator re-runs the same engine with modified inputs: shifting customer invoice dates (optionally for one named
customer), shifting supplier payments, scaling revenue or expenses, adding a purchase, financing with monthly cost,
collecting overdue invoices within N days. It returns baseline vs scenario curves and the change in the lowest point.</p>
{img("08_simulator_result","Simulator with a scenario result")}
<h3>7.4 Rescue plans</h3>
<p>Candidate strategies (collect overdue invoices, negotiate supplier extensions, bridge with a short-term credit line,
labeled illustrative) are each simulated and ranked by the improvement of the lowest cash point. On Northstar Studio the
measured improvements are +EUR 7,150, +EUR 4,560 and +EUR 4,372.</p>
<h3>7.5 Alerts</h3>
<p>Summary and forecast requests evaluate rules (projected shortfall, overdue invoices) and create notifications with a
7-day deduplication window.</p>

<h2>8. Agentic AI architecture</h2>
<table><tr><th>Agent</th><th>Tools</th><th>Role</th></tr>
<tr><td>Financial Analyst</td><td>get_financial_summary, list_invoices</td><td>Current position</td></tr>
<tr><td>Forecast Analyst</td><td>forecast_cash</td><td>Future balance and shortfall</td></tr>
<tr><td>Strategic Planner</td><td>simulate_scenario, get_rescue_plans</td><td>What-if and strategy ranking</td></tr>
<tr><td>Workflow Coordinator</td><td>propose_invoice_reminder</td><td>Creates PENDING_APPROVAL actions only</td></tr>
<tr><td>Risk Guardian</td><td>role checks, final review</td><td>Denies tools above the user role</td></tr></table>
<p>The orchestrator uses OpenAI tool calling (gpt-4o-mini) with at most 6 steps. Tool results are deterministic engine
outputs. If the model is unavailable, a deterministic intent router answers from the same tools. The trace (agents and
tools used) is returned and displayed under each answer.</p>
{img("12_copilot_agent_trace","Copilot answer for a late-payment question, with agent trace (Strategic Planner, Risk Guardian)")}

<h2>9. Human-in-the-loop execution</h2>
<p>Actions follow PROPOSE (ACCOUNTANT or above) then APPROVE (role checked) then EXECUTE. Executing an unapproved action
returns HTTP 409; executing twice returns 409 (idempotent). If no n8n webhook is configured, the action is recorded as
QUEUED with an explicit message and is never reported as SENT.</p>
{img("10_decisions","Decision Center")}

<h2>10. Security design</h2>
<p>JWT authentication, bcrypt password hashes, email verification codes stored as HMAC hashes with 10-minute expiry,
5 attempts and 60 s resend cooldown, organization-scoped queries, role hierarchy, admin-only console, audit log,
server-side Google Sheets URL reconstruction (SSRF guard), secrets only in environment variables. Full detail and known
risks are in PULSE_Security_Self_Assessment.pdf.</p>

<h2>11. Integrations</h2>
<p>Supabase Postgres (session pooler), OpenAI API, SMTP (Gmail app password) for verification codes, n8n workflows in
<code>n8n/workflows/</code> (verification code, invoice reminder, cash risk alert, financial digest, approved recovery plan)
protected by a header secret.</p>

<h2>12. Testing</h2>
<p>87 automated backend tests (pytest, SQLite) pass: auth and OTP, onboarding, imports, financial summary, forecast and
simulation regressions, rescue plans, agents and tool permissions, actions workflow, notifications, admin. Output in
<code>execution_evidence/test_output.txt</code>. Frontend: TypeScript strict build passes.</p>

<h2>13. Deployment</h2>
<pre>cd backend && python -m venv .venv && .venv/Scripts/activate
pip install -r requirements.txt
cp .env.example .env      # DATABASE_URL, JWT secret, optional OPENAI_API_KEY, SMTP_*
alembic upgrade head
python scripts/seed_demo.py         # Northstar Studio company
uvicorn app.main:app --reload       # http://127.0.0.1:8000/docs
cd ../frontend && npm install && npm run dev
cd backend && pytest -q</pre>

<h2>14. Limitations and roadmap</h2>
<ul><li>Forecast is deterministic and rule-based; no ML model is trained yet.</li>
<li>External sending depends on n8n/SMTP configuration; otherwise actions stay QUEUED.</li>
<li>No direct bank or ERP connector yet (Odoo connector planned); data comes from file, Google Sheets or manual entry.</li>
<li>No rate limiting, MFA, or field-level encryption yet (see security report).</li>
<li>Credit line strategy uses illustrative terms, not a lender offer.</li></ul>
</body></html>"""


def security_doc():
    rows = [
        ("Authentication", "JWT bearer tokens (python-jose), bcrypt password hashing", "ok"),
        ("Email verification", "6-digit code, HMAC-hashed at rest, 10 min TTL, 5 attempts, 60 s resend cooldown", "ok"),
        ("Account approval", "Accounts approved after verified email code; admins can suspend; reviews stored in AccountReview", "ok"),
        ("Tenant isolation", "Every financial query filtered by organization_id after membership check", "ok"),
        ("Role-based access", "OWNER > FINANCE_MANAGER > ACCOUNTANT > VIEWER enforced by require_role", "ok"),
        ("Agent tool permissions", "Risk Guardian denies propose_invoice_reminder below ACCOUNTANT (ToolDenied)", "ok"),
        ("Human approval", "PROPOSE, APPROVE, EXECUTE; unapproved or repeated execution returns 409", "ok"),
        ("Honest status", "Actions without a configured webhook are QUEUED, never reported SENT", "ok"),
        ("Audit trail", "AuditLog for approvals, executions and admin decisions; admin activity log page", "ok"),
        ("SSRF guard", "Google Sheets export URL rebuilt server-side; only docs.google.com fetched", "ok"),
        ("Secrets", "Keys only in backend/.env (gitignored); not in repo or submission ZIP", "ok"),
        ("Webhook auth", "n8n webhooks require a shared header secret", "ok"),
        ("LLM grounding", "Numbers come from deterministic tools; deterministic fallback if the LLM fails", "ok"),
        ("Input validation", "Pydantic schemas, Decimal money, import column validation", "part"),
        ("Transport security", "HTTPS depends on deployment; local dev runs over HTTP", "part"),
        ("CORS", "Restricted to configured frontend origins", "part"),
        ("Rate limiting / brute force", "Only OTP attempt limits; no global login rate limiting", "plan"),
        ("MFA for owners/admins", "Not implemented", "plan"),
        ("Encryption of financial fields at rest", "Relies on Supabase disk encryption only", "plan"),
        ("Token revocation / refresh", "Short-lived JWT only; no revocation list", "plan"),
        ("Dependency and SAST scanning in CI", "Not set up", "plan"),
    ]
    label = {"ok": "IMPLEMENTED", "part": "PARTIAL", "plan": "PLANNED"}
    trs = "".join(f'<tr><td>{a}</td><td>{b}</td><td class="{c}">{label[c]}</td></tr>' for a, b, c in rows)
    tiers = """<table><tr><th>Capability</th><th>VIEWER</th><th>ACCOUNTANT</th><th>FINANCE_MANAGER</th><th>OWNER</th><th>Platform admin</th></tr>
<tr><td>View dashboard, forecast, copilot answers</td><td>Yes</td><td>Yes</td><td>Yes</td><td>Yes</td><td>No (admin console only)</td></tr>
<tr><td>Run simulations</td><td>Yes</td><td>Yes</td><td>Yes</td><td>Yes</td><td>No</td></tr>
<tr><td>Import data</td><td>No</td><td>Yes</td><td>Yes</td><td>Yes</td><td>No</td></tr>
<tr><td>Propose actions (incl. via agents)</td><td>No</td><td>Yes</td><td>Yes</td><td>Yes</td><td>No</td></tr>
<tr><td>Approve and execute actions</td><td>No</td><td>No</td><td>Yes</td><td>Yes</td><td>No</td></tr>
<tr><td>Manage members and organization</td><td>No</td><td>No</td><td>No</td><td>Yes</td><td>No</td></tr>
<tr><td>Review / suspend accounts, view audit</td><td>No</td><td>No</td><td>No</td><td>No</td><td>Yes</td></tr></table>
<p class="note">Approval and execution thresholds are as enforced in <code>backend/app/api/actions.py</code> and
<code>app/api/deps.py</code>. Check the code for the exact minimum role of each endpoint.</p>"""
    risks = """<table><tr><th>#</th><th>Known risk</th><th>Impact</th><th>Mitigation status</th></tr>
<tr><td>R1</td><td>Credential stuffing on /auth/login</td><td>Account takeover</td><td class="plan">PLANNED: rate limiting, MFA</td></tr>
<tr><td>R2</td><td>Prompt injection through imported descriptions or chat</td><td>Misleading answer</td><td class="part">PARTIAL: tools are read-only or create PENDING_APPROVAL only; numbers from engines</td></tr>
<tr><td>R3</td><td>Sensitive financial data sent to the LLM provider</td><td>Privacy</td><td class="part">PARTIAL: only tool results for the user's org; provider can be disabled (deterministic mode)</td></tr>
<tr><td>R4</td><td>Leaked JWT used until expiry</td><td>Session hijack</td><td class="plan">PLANNED: refresh tokens and revocation</td></tr>
<tr><td>R5</td><td>Misconfigured webhook secret</td><td>Spoofed workflow calls</td><td class="part">PARTIAL: header secret required; rotation manual</td></tr>
<tr><td>R6</td><td>Forecast error on irregular businesses</td><td>Wrong decision</td><td class="ok">MITIGATED: known vs uncertain split, human approval, rescue plans labeled as simulations</td></tr>
<tr><td>R7</td><td>Malformed or huge import files</td><td>Bad data, resource use</td><td class="part">PARTIAL: preview, mapping, validation; size limits to harden</td></tr>
<tr><td>R8</td><td>Dependency vulnerabilities</td><td>Varies</td><td class="plan">PLANNED: pip-audit / npm audit in CI</td></tr></table>"""
    return f"""<html><head><meta charset="utf-8"><style>{CSS}</style></head><body>
<div class="cover"><h1>PULSE AI</h1><p>Security Self-Assessment Report</p><p>FinTechathon 2026, Topic E</p></div>
<h2>1. Scope</h2><p>Backend API (FastAPI), frontend (React), Supabase Postgres, AI agents, n8n workflows. This assessment
is a self-review by the team; no external penetration test was performed.</p>
<h2>2. Permission-tier implementation</h2>{tiers}
<h2>3. Controls</h2><table><tr><th>Control</th><th>Implementation</th><th>Status</th></tr>{trs}</table>
<h2>4. Verified behaviour</h2><p>Captured against the running backend (see execution_evidence/api_core_scenario.md):
unauthenticated request returns 401; executing an unapproved action returns 409; approval by OWNER returns 200;
execution without webhook returns QUEUED; repeated execution returns 409. Agent permission denial is covered by
automated tests (87 passing).</p>
<h2>5. Known-risk list</h2>{risks}
<h2>6. Data handling</h2><p>Submission evidence uses the Northstar Studio company only. Identifiers, emails and tokens
are masked in logs and screenshots. No .env file or credential is included in the submission package.</p>
{img("19_admin_audit","Admin activity log (personal emails masked)")}
</body></html>"""


SLIDES = [
    ("The Invisible Financial Crisis", "A profitable SME can still run out of cash.\nThe shortfall is visible weeks ahead, but nobody is looking.", "01_landing",
     "Open with the problem: cash, not profit, kills small businesses. Late payments and clustered bills create a gap that spreadsheets hide."),
    ("Why SMEs Struggle", "No finance team\nData scattered across bank exports and sheets\nTools describe the past, not the next 60 days", None,
     "SMEs lack the time and tooling to forecast. Explain the three gaps: visibility, foresight, action."),
    ("Introducing PULSE AI", "A cash-flow copilot that predicts, simulates, and proposes,\nwhile a human keeps the final decision.", "04_dashboard",
     "Show the dashboard of Northstar Studio: current cash EUR 7,020, shortfall predicted on 7 Nov 2026."),
    ("The Financial Digital Twin", "Replay the business under any scenario:\ncustomer delays, supplier terms, financing, revenue change.", "08_simulator_result",
     "The simulator reuses the exact forecast engine with modified inputs, so baseline and scenario are comparable."),
    ("Predicting Financial Risk", "Deterministic day-by-day forecast\nKnown flows plus uncertain baseline\nLowest point, shortfall date, risk level", "06_forecast",
     "Every number is computed, not generated. Lowest point -EUR 2,308 on 7 Nov 2026 for Northstar Studio."),
    ("From Risk to Rescue", "Strategies simulated and ranked by impact:\nCollect overdue +EUR 7,150\nCredit line +EUR 4,560 (illustrative)\nSupplier extension +EUR 4,372", "05_dashboard_full",
     "Rescue plans are simulations, not guarantees. The credit line uses illustrative terms."),
    ("Agentic AI Architecture", "Financial Analyst, Forecast Analyst, Strategic Planner,\nWorkflow Coordinator, Risk Guardian\nOpenAI tool calling with deterministic fallback", "12_copilot_agent_trace",
     "The model picks tools; tools return engine numbers. The trace under each answer shows which agents acted."),
    ("Live Demo and Execution Evidence", "Propose, approve, execute, recorded in the audit log\nUnapproved execution refused (409)\nNo webhook: QUEUED, never falsely SENT", "10_decisions",
     "Walk through execution_evidence/api_core_scenario.md. 87 automated tests pass."),
    ("Security and Human Control", "JWT, bcrypt, hashed email codes\nRole tiers from VIEWER to OWNER\nAdmin console with account review and audit", "14_admin_overview",
     "Refer to the security self-assessment: implemented, partial and planned controls are listed separately."),
    ("Impact and Future Vision", "Weeks of warning instead of a surprise overdraft\nNext: ERP connectors (Odoo), bank feeds, MFA,\nlearned forecasting on top of the deterministic core", None,
     "Close with the vision. Do not quote impact metrics we have not measured."),
]


def build_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    navy, cyan, teal, white = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0x22, 0xD3, 0xEE), RGBColor(0x14, 0xB8, 0xA6), RGBColor(255, 255, 255)
    for i, (title, body, shot, notes) in enumerate(SLIDES, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = navy
        bar = s.shapes.add_shape(1, Inches(0.6), Inches(0.55), Inches(0.12), Inches(0.9)); bar.fill.solid(); bar.fill.fore_color.rgb = teal; bar.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.9), Inches(0.45), Inches(11.5), Inches(1.1)).text_frame
        tb.text = title; p = tb.paragraphs[0]; p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = white; p.font.name = "Poppins"
        w = Inches(4.6) if shot else Inches(11.5)
        bt = s.shapes.add_textbox(Inches(0.9), Inches(1.9), w, Inches(4.8)).text_frame; bt.word_wrap = True
        for j, line in enumerate(body.split("\n")):
            para = bt.paragraphs[0] if j == 0 else bt.add_paragraph()
            para.text = line; para.font.size = Pt(20 if shot else 26); para.font.color.rgb = cyan if j == 0 else white; para.font.name = "Inter"; para.space_after = Pt(12)
        if shot:
            s.shapes.add_picture(str(SHOTS / f"{shot}.png"), Inches(5.8), Inches(1.8), width=Inches(7.0))
        ft = s.shapes.add_textbox(Inches(0.9), Inches(6.9), Inches(11.5), Inches(0.4)).text_frame
        ft.text = f"PULSE AI  |  FinTechathon 2026, Topic E  |  {i}/10"; ft.paragraphs[0].font.size = Pt(11); ft.paragraphs[0].font.color.rgb = teal
        s.notes_slide.notes_text_frame.text = notes
    prs.save(SUB / "PULSE_Presentation.pptx")


def deck_html():
    parts = []
    for i, (title, body, shot, _) in enumerate(SLIDES, 1):
        lines = body.split("\n")
        txt = f'<p class="lead">{lines[0]}</p>' + "".join(f"<p>{l}</p>" for l in lines[1:])
        pic = ""
        if shot:
            data = base64.b64encode((SHOTS / f"{shot}.png").read_bytes()).decode()
            pic = f'<img src="data:image/png;base64,{data}"/>'
        parts.append(f'<section><h1>{title}</h1><div class="row"><div class="t {"full" if not shot else ""}">{txt}</div>{pic}</div><footer>PULSE AI | FinTechathon 2026, Topic E | {i}/10</footer></section>')
    css = f"""@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600&family=Poppins:wght@700&display=swap');
@page{{size:13.333in 7.5in;margin:0}} body{{margin:0}}
section{{width:13.333in;height:7.5in;background:{NAVY};color:white;padding:.5in .8in;box-sizing:border-box;page-break-after:always;position:relative;font-family:Inter,Arial}}
h1{{font-family:Poppins,Arial;font-size:36pt;margin:0 0 .4in;border-left:10px solid {TEAL};padding-left:18px}}
.row{{display:flex;gap:.5in;align-items:flex-start}} .t{{width:4.6in;font-size:19pt}} .t.full{{width:11in;font-size:26pt}}
.lead{{color:{CYAN};font-weight:600}} img{{width:7in;border-radius:10px;border:2px solid {TEAL}}}
footer{{position:absolute;bottom:.3in;left:.8in;color:{TEAL};font-size:11pt}}"""
    return f'<html><head><meta charset="utf-8"><style>{css}</style></head><body>{"".join(parts)}</body></html>'


def main():
    build_pptx()
    with sync_playwright() as p:
        br = p.chromium.launch()
        pg = br.new_page()
        for name, html, opts in [
            ("PULSE_Technical_Documentation.pdf", tech_doc(), {"format": "A4", "margin": {"top": "18mm", "bottom": "18mm", "left": "16mm", "right": "16mm"}}),
            ("PULSE_Security_Self_Assessment.pdf", security_doc(), {"format": "A4", "margin": {"top": "18mm", "bottom": "18mm", "left": "16mm", "right": "16mm"}}),
            ("PULSE_Presentation.pdf", deck_html(), {"width": "13.333in", "height": "7.5in"}),
        ]:
            pg.set_content(html, wait_until="networkidle")
            pg.pdf(path=str(SUB / name), print_background=True, **opts)
            print("built", name)
        br.close()
    print("built PULSE_Presentation.pptx")


if __name__ == "__main__":
    main()
